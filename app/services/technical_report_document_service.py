from __future__ import annotations

from pathlib import Path
import zipfile

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.text import WD_COLOR_INDEX
from docx.shared import Inches, Pt, RGBColor


MISSING_MARKER = "PENDENTE DE REVISÃO - informação não registrada no chamado"


def _replace_paragraph_tokens(paragraph, context: dict[str, object], missing: set[str]) -> None:
    for run in paragraph.runs:
        for token, value in context.items():
            marker = "{{ " + token + " }}"
            if marker not in run.text:
                continue
            text = str(value or "")
            is_missing = token in missing
            if not text:
                text = MISSING_MARKER
                is_missing = True
            run.text = run.text.replace(marker, text)
            if is_missing:
                run.font.highlight_color = WD_COLOR_INDEX.YELLOW
                run.font.bold = True
                run.font.color.rgb = RGBColor(128, 80, 0)


def _replace_table_tokens(table, context: dict[str, object], missing: set[str]) -> None:
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                _replace_paragraph_tokens(paragraph, context, missing)
            for nested in cell.tables:
                _replace_table_tokens(nested, context, missing)


def _style_document(document) -> None:
    for style_name in ("Normal", "Title", "Subtitle", "Heading 1", "Heading 2"):
        style = document.styles[style_name]
        style.font.name = "Aptos"
        style.font.color.rgb = RGBColor(0, 0, 0)
    document.styles["Normal"].font.size = Pt(10.5)
    document.styles["Title"].font.size = Pt(22)
    document.styles["Title"].font.bold = True
    document.styles["Heading 1"].font.size = Pt(13)
    document.styles["Heading 1"].font.bold = True
    document.styles["Heading 2"].font.size = Pt(11)
    for section in document.sections:
        section.top_margin = Inches(0.7)
        section.bottom_margin = Inches(0.65)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)
        for paragraph in section.footer.paragraphs:
            _replace_paragraph_tokens(paragraph, {}, set())


def render_technical_report(
    template_path: Path,
    context: dict[str, object],
    missing_fields: list[str],
    output_path: Path,
) -> Path:
    if not template_path.is_file():
        raise FileNotFoundError("O modelo de relatório técnico não está disponível.")
    output_path.parent.mkdir(parents=True, exist_ok=True)
    document = Document(str(template_path))
    _style_document(document)
    missing = set(missing_fields)
    for paragraph in document.paragraphs:
        _replace_paragraph_tokens(paragraph, context, missing)
        if paragraph.style.name == "Title":
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for table in document.tables:
        _replace_table_tokens(table, context, missing)
    for section in document.sections:
        for paragraph in section.header.paragraphs:
            _replace_paragraph_tokens(paragraph, context, missing)
        for paragraph in section.footer.paragraphs:
            _replace_paragraph_tokens(paragraph, context, missing)
    document.save(str(output_path))
    with zipfile.ZipFile(output_path) as archive:
        if "word/document.xml" not in archive.namelist():
            raise ValueError("O documento gerado não é um pacote DOCX válido.")
    return output_path
