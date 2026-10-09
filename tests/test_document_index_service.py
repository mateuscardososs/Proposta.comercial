from __future__ import annotations

import subprocess
import tempfile
import time
from datetime import date, datetime, timezone
from pathlib import Path
from types import SimpleNamespace

import pytest
from docx import Document

from app.assistant.service import AssistantService
from app.models import (
    AssistantAction,
    AssistantConversation,
    Client,
    Proposal,
    ServiceCall,
    ServiceTechnicalReport,
    User,
)
from app.services.document_index_service import (
    LocalOCRUnavailable,
    local_ocr_for_platform,
    local_tesseract_ocr,
    reindex_registered_documents,
)
from app.services.document_search_service import search_document_index


def _pdf(text: str) -> bytes:
    escaped = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
    stream = f"BT /F1 12 Tf 72 720 Td ({escaped}) Tj ET".encode("latin-1")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    payload = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for index, obj in enumerate(objects, start=1):
        offsets.append(len(payload))
        payload.extend(f"{index} 0 obj\n".encode())
        payload.extend(obj + b"\nendobj\n")
    xref_offset = len(payload)
    payload.extend(f"xref\n0 {len(offsets)}\n".encode())
    payload.extend(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        payload.extend(f"{offset:010d} 00000 n \n".encode())
    payload.extend(
        f"trailer\n<< /Size {len(offsets)} /Root 1 0 R >>\nstartxref\n{xref_offset}\n%%EOF".encode()
    )
    return bytes(payload)


def _proposal(db, root: Path) -> Proposal:
    root.mkdir(parents=True, exist_ok=True)
    client = Client(razao_social="Cliente Sintético")
    user = User(nome="Usuário sintético", email="index@example.invalid", senha_hash="not-a-secret")
    db.add_all([client, user])
    db.flush()
    docx = Document()
    docx.add_heading("Dados do equipamento", level=1)
    docx.add_paragraph("Número de série sintético: BAL-4931.")
    docx_path = root / "proposta-indexada.docx"
    docx.save(docx_path)
    pdf_path = root / "proposta-indexada.pdf"
    pdf_path.write_bytes(_pdf("Valor total da proposta: R$ 8.450,00"))
    row = Proposal(
        numero=711,
        revisao="00",
        client_id=client.id,
        user_id=user.id,
        origem="upload_externo",
        docx_path=docx_path.name,
        pdf_path=pdf_path.name,
    )
    db.add(row)
    db.commit()
    return row


def test_report_word_does_not_route_an_ordinary_task_request_to_document_search():
    assert not AssistantService._direct_document_query(
        "Preciso que preparem uma atividade de revisar relatório"
    )
    assert AssistantService._direct_document_query(
        "O que consta no relatório técnico sobre a verificação?"
    )


def test_persists_searchable_pdf_page_and_docx_section(db, tmp_path):
    _proposal(db, tmp_path)

    first = reindex_registered_documents(db, output_dir=tmp_path)
    db.commit()

    assert first.indexed == 2
    pdf_result = search_document_index(db, query="valor total proposta")
    docx_result = search_document_index(db, query="número série equipamento")
    assert pdf_result.evidence[0].page == 1
    assert "8.450" in pdf_result.evidence[0].excerpt
    assert docx_result.evidence[0].section == "Dados do equipamento"
    assert "BAL-4931" in docx_result.evidence[0].excerpt


def test_reindex_is_idempotent_and_replacement_removes_stale_text(db, tmp_path):
    proposal = _proposal(db, tmp_path)
    reindex_registered_documents(db, output_dir=tmp_path)
    db.commit()
    count_before = search_document_index(db, query="valor total").registered_documents

    again = reindex_registered_documents(db, output_dir=tmp_path)
    db.commit()
    assert again.unchanged == 2
    assert search_document_index(db, query="valor total").registered_documents == count_before

    (tmp_path / proposal.pdf_path).write_bytes(_pdf("Valor revisado da proposta: R$ 9.700,00"))
    replaced = reindex_registered_documents(db, output_dir=tmp_path)
    db.commit()
    assert replaced.updated == 1
    assert not search_document_index(db, query="8.450").evidence
    assert "9.700" in search_document_index(db, query="valor revisado").evidence[0].excerpt


def test_deleted_registered_file_is_removed_from_index_and_not_claimed_as_empty(db, tmp_path):
    proposal = _proposal(db, tmp_path)
    reindex_registered_documents(db, output_dir=tmp_path)
    db.commit()
    (tmp_path / proposal.pdf_path).unlink()

    result = reindex_registered_documents(db, output_dir=tmp_path)
    db.commit()
    search = search_document_index(db, query="valor total")

    assert result.missing == 1
    assert not search.evidence
    assert search.unreadable_documents == 1


def test_unregistered_source_is_removed_from_index_on_reconciliation(db, tmp_path):
    proposal = _proposal(db, tmp_path)
    reindex_registered_documents(db, output_dir=tmp_path)
    db.commit()
    proposal.pdf_path = ""
    proposal.docx_path = ""
    db.commit()

    result = reindex_registered_documents(db, output_dir=tmp_path)
    db.commit()

    assert result.removed == 2
    assert search_document_index(db, query="valor total").registered_documents == 0


def test_scanned_portuguese_pdf_uses_local_vision_ocr_and_is_searchable(db, tmp_path):
    image_module = pytest.importorskip("PIL.Image")
    draw_module = pytest.importorskip("PIL.ImageDraw")
    font_module = pytest.importorskip("PIL.ImageFont")
    root = tmp_path / "scan"
    root.mkdir()
    font = font_module.truetype("/System/Library/Fonts/Supplemental/Arial.ttf", 42)
    image = image_module.new("RGB", (1500, 500), "white")
    draw_module.Draw(image).text((65, 150), "Inspeção técnica da balança", fill="black", font=font)
    scanned_path = root / "proposta-escaneada.pdf"
    image.save(scanned_path, "PDF", resolution=180)
    client = Client(razao_social="Cliente OCR Sintético")
    user = User(nome="Usuário OCR", email="ocr@example.invalid", senha_hash="not-a-secret")
    db.add_all([client, user])
    db.flush()
    db.add(Proposal(
        numero=712,
        revisao="00",
        client_id=client.id,
        user_id=user.id,
        origem="upload_externo",
        pdf_path=scanned_path.name,
    ))
    db.commit()

    result = reindex_registered_documents(db, output_dir=root)
    db.commit()
    search = search_document_index(db, query="inspeção técnica balança")

    assert result.ocr_pages == 1
    assert search.evidence
    assert "Inspeção" in search.evidence[0].excerpt or "Inspecao" in search.evidence[0].excerpt
    assert search.evidence[0].page == 1


def test_technical_report_files_are_registered_search_sources(db, tmp_path):
    client = Client(razao_social="Cliente do Chamado")
    db.add(client)
    db.flush()
    conversation = AssistantConversation()
    db.add(conversation)
    db.flush()
    action = AssistantAction(
        conversation_id=conversation.id,
        request_id="report-index-action",
        confirmation_token_hash="a" * 64,
        action_type="generate_service_report",
        status="executed",
        arguments_json={},
        result_json={},
    )
    call = ServiceCall(client_id=client.id, summary="Inspeção sintética", opened_on=date(2026, 10, 1))
    db.add_all([action, call])
    db.flush()
    payload = Document()
    payload.add_heading("Resultado da verificação", level=1)
    payload.add_paragraph("A balança ficou estável após calibração.")
    docx_path = tmp_path / "relatorio-sintetico.docx"
    payload.save(docx_path)
    pdf_path = tmp_path / "relatorio-sintetico.pdf"
    pdf_path.write_bytes(_pdf("Resultado da verificação: balança estável após calibração"))
    db.add(ServiceTechnicalReport(
        service_call_id=call.id,
        assistant_action_id=action.id,
        document_event_id=None,
        idempotency_key="report-index-1",
        source_fingerprint="b" * 64,
        source_event_ids=[],
        client_snapshot_json={},
        fields_json={},
        source_fields_json={},
        manual_overrides_json=[],
        missing_fields_json=[],
        docx_path=docx_path.name,
        pdf_path=pdf_path.name,
        confirmed_at=datetime(2026, 10, 1, tzinfo=timezone.utc),
    ))
    db.commit()

    result = reindex_registered_documents(db, output_dir=tmp_path)
    db.commit()
    search = search_document_index(db, query="resultado verificação balança estável")
    reply = AssistantService(db, None, output_dir=tmp_path).handle_message(
        message="O relatório técnico do chamado fala do resultado da verificação?",
        request_id="report-index-assistant-query",
    )

    assert result.indexed == 2
    assert any(evidence.document_name == "relatorio-sintetico.docx" for evidence in search.evidence)
    assert any(evidence.document_name == "relatorio-sintetico.pdf" for evidence in search.evidence)
    assert "Resultado da verifica" in reply.message
    assert "relatorio-sintetico.pdf" in reply.message


def test_unavailable_ocr_marks_scanned_pdf_unsearchable(db, tmp_path):
    image_module = pytest.importorskip("PIL.Image")
    image = image_module.new("RGB", (300, 200), "white")
    path = tmp_path / "sem-ocr.pdf"
    image.save(path, "PDF")
    client = Client(razao_social="Cliente OCR indisponível")
    user = User(nome="Usuário OCR", email="ocr-down@example.invalid", senha_hash="not-a-secret")
    db.add_all([client, user])
    db.flush()
    proposal = Proposal(
        numero=713,
        revisao="00",
        client_id=client.id,
        user_id=user.id,
        origem="upload_externo",
        pdf_path=path.name,
    )
    db.add(proposal)
    db.commit()

    result = reindex_registered_documents(db, output_dir=tmp_path, ocr_adapter=lambda _path: None)
    db.commit()
    search = search_document_index(db, query="conteúdo escaneado")

    assert result.unsearchable == 1
    assert search.unreadable_documents == 1
    assert not search.evidence


def test_assistant_does_not_claim_absence_when_local_ocr_is_unavailable(db, tmp_path, monkeypatch):
    from PIL import Image

    import app.services.document_index_service as index_service

    scan = Image.new("RGB", (500, 250), "white")
    scan_path = tmp_path / "escaneado.pdf"
    scan.save(scan_path, "PDF")
    client = Client(razao_social="Cliente OCR indisponível")
    user = User(nome="Usuário OCR", email="ocr-chat@example.invalid", senha_hash="not-a-secret")
    db.add_all([client, user])
    db.flush()
    db.add(Proposal(
        numero=714,
        revisao="00",
        client_id=client.id,
        user_id=user.id,
        origem="upload_externo",
        pdf_path=scan_path.name,
    ))
    db.commit()
    monkeypatch.setattr(index_service, "local_vision_ocr", lambda _path: None)

    reply = AssistantService(db, None, output_dir=tmp_path).handle_message(
        message="O que consta no documento 714 sobre calibração?",
        request_id="ocr-unavailable-assistant",
    )

    assert "não consegui consultar" in reply.message.lower()
    assert "ocr local não está disponível" in reply.message.lower()
    assert "não encontrei evidência suficiente" not in reply.message.lower()


def test_platform_ocr_selection_preserves_apple_vision_and_uses_tesseract_on_windows(monkeypatch, tmp_path):
    import app.services.document_index_service as index_service

    calls: list[str] = []
    monkeypatch.setattr(index_service, "local_vision_ocr", lambda _path: calls.append("vision") or ["mac"])
    monkeypatch.setattr(index_service, "local_tesseract_ocr", lambda _path: calls.append("tesseract") or ["windows"])

    assert local_ocr_for_platform(tmp_path / "scan.pdf", system="Darwin") == ["mac"]
    assert local_ocr_for_platform(tmp_path / "scan.pdf", system="Windows") == ["windows"]
    assert local_ocr_for_platform(tmp_path / "scan.pdf", system="Linux") is None
    assert calls == ["vision", "tesseract"]


def test_windows_tesseract_uses_portuguese_and_returns_page_text(monkeypatch, tmp_path):
    image_page = b"synthetic png bytes"
    calls: list[list[str]] = []
    temp_dirs: list[Path] = []

    def fake_temp_dir(**_kwargs):
        from tempfile import TemporaryDirectory

        temp = TemporaryDirectory(dir=tmp_path)
        temp_dirs.append(Path(temp.name))
        return temp

    def fake_render(_pdf_path):
        return [image_page, image_page]

    def fake_run(args, **kwargs):
        calls.append(args)
        assert kwargs["shell"] is False
        if "--list-langs" in args:
            return SimpleNamespace(returncode=0, stdout=b"List of available languages (2):\neng\npor\n", stderr=b"")
        return SimpleNamespace(returncode=0, stdout=b"Texto OCR em portugu\xc3\xaas", stderr=b"")

    pages = local_tesseract_ocr(
        tmp_path / "scan.pdf",
        executable="C:/Program Files/Tesseract-OCR/tesseract.exe",
        data_dir="C:/Program Files/Tesseract-OCR/tessdata",
        which=lambda _name: "C:/Program Files/Tesseract-OCR/tesseract.exe",
        runner=fake_run,
        renderer=fake_render,
        temporary_directory=fake_temp_dir,
    )

    assert pages == ["Texto OCR em português", "Texto OCR em português"]
    assert any("por" in call for call in calls[1:])
    assert any("--tessdata-dir" in call for call in calls)
    assert all(not directory.exists() for directory in temp_dirs)


def test_windows_ocr_reports_missing_executable_and_portuguese_data(tmp_path):
    with pytest.raises(LocalOCRUnavailable, match="tesseract_missing"):
        local_tesseract_ocr(tmp_path / "scan.pdf", executable=None)

    def no_portuguese(args, **kwargs):
        return SimpleNamespace(returncode=0, stdout=b"List of available languages (1):\neng\n", stderr=b"")

    with pytest.raises(LocalOCRUnavailable, match="ocr_language_missing"):
        local_tesseract_ocr(
            tmp_path / "scan.pdf",
            executable="tesseract.exe",
            which=lambda name: name,
            runner=no_portuguese,
            renderer=lambda _path: [],
        )


def test_windows_ocr_conversion_failure_cleans_temporary_directory(tmp_path):
    temp_dirs: list[Path] = []

    def tracked_temp_dir(**_kwargs):
        from tempfile import TemporaryDirectory

        temp = TemporaryDirectory(dir=tmp_path)
        temp_dirs.append(Path(temp.name))
        return temp

    def failed_render(_pdf_path):
        raise RuntimeError("synthetic conversion failure")

    with pytest.raises(LocalOCRUnavailable, match="ocr_render_failed"):
        local_tesseract_ocr(
            tmp_path / "scan.pdf",
            executable="tesseract.exe",
            which=lambda name: name,
            runner=lambda *_args, **_kwargs: SimpleNamespace(
                returncode=0, stdout=b"por", stderr=b""
            ),
            renderer=failed_render,
            temporary_directory=tracked_temp_dir,
        )

    assert temp_dirs and all(not directory.exists() for directory in temp_dirs)


@pytest.mark.parametrize("failure,code", [
    ("deadline", "ocr_tesseract_timeout"),
    ("subprocess", "ocr_tesseract_failed"),
    ("limit", "ocr_page_limit"),
])
def test_local_ocr_failure_codes_close_renderer_and_temporary_directory(
    monkeypatch, tmp_path, failure, code
):
    closed = []
    directories = []
    calls = []
    ticks = iter([0, 91] if failure == "deadline" else [0] + [1] * 501)
    monkeypatch.setattr(time, "monotonic", lambda: next(ticks))

    def renderer(_path):
        try:
            for _ in range(501 if failure == "limit" else 1):
                yield b"synthetic PNG"
        finally:
            closed.append(True)

    def directory(**kwargs):
        assert kwargs == {"prefix": "adbalancas-ocr-"}
        temporary = tempfile.TemporaryDirectory(dir=tmp_path)
        directories.append(Path(temporary.name))
        return temporary

    def runner(args, **kwargs):
        calls.append((args, kwargs))
        assert kwargs["shell"] is False
        if "--list-langs" in args:
            assert kwargs["timeout"] == 10
            return SimpleNamespace(returncode=0, stdout="por\n")
        assert kwargs["timeout"] == 30
        if failure == "subprocess":
            raise subprocess.TimeoutExpired(args, kwargs["timeout"])
        return SimpleNamespace(returncode=0, stdout=" texto \n")

    with pytest.raises(LocalOCRUnavailable) as caught:
        local_tesseract_ocr(
            tmp_path / "synthetic.pdf", executable="tesseract", data_dir=None,
            which=lambda name: name, runner=runner, renderer=renderer,
            temporary_directory=directory,
        )
    assert caught.value.code == code
    assert str(caught.value) == code
    assert closed == [True]
    assert directories and all(not path.exists() for path in directories)
    assert len(calls) == (501 if failure == "limit" else 2 if failure == "subprocess" else 1)


def test_local_ocr_extraction_preserves_exception_and_callable_identity():
    from app.services import document_index_service, local_pdf_ocr

    assert document_index_service.LocalOCRUnavailable is local_pdf_ocr.LocalOCRUnavailable
    assert document_index_service.local_tesseract_ocr is local_pdf_ocr.local_tesseract_ocr
    assert document_index_service.local_vision_ocr is local_pdf_ocr.local_vision_ocr
    assert document_index_service.OCR_TIMEOUT_SECONDS == 90
    assert document_index_service.TESSERACT_PAGE_TIMEOUT_SECONDS == 30
    assert document_index_service.TESSERACT_MAX_PAGES == 500


def test_missing_portuguese_language_is_reported_in_document_search(db, tmp_path):
    from PIL import Image

    from app.services.document_index_service import LocalOCRUnavailable

    scanned = tmp_path / "idioma-ausente.pdf"
    Image.new("RGB", (300, 200), "white").save(scanned, "PDF")
    client = Client(razao_social="Cliente idioma OCR sintético")
    user = User(nome="Usuário idioma", email="ocr-lang@example.invalid", senha_hash="not-a-secret")
    db.add_all([client, user])
    db.flush()
    proposal = Proposal(
        numero=716,
        revisao="00",
        client_id=client.id,
        user_id=user.id,
        origem="upload_externo",
        pdf_path=scanned.name,
    )
    db.add(proposal)
    db.commit()

    def missing_language(_path):
        raise LocalOCRUnavailable("ocr_language_missing")

    result = reindex_registered_documents(db, output_dir=tmp_path, ocr_adapter=missing_language)
    db.commit()
    search = search_document_index(db, query="informação sintética")

    assert result.unsearchable == 1
    assert search.ocr_unavailable_documents == 1
    assert search.ocr_unavailable_reasons == (
        "dados do idioma português (por) ausentes no tessdata",
    )


def test_assistant_explains_missing_windows_tesseract_without_claiming_empty_documents(
    db, tmp_path, monkeypatch
):
    from PIL import Image

    import app.services.document_index_service as index_service

    scanned = tmp_path / "sem-tesseract.pdf"
    Image.new("RGB", (300, 200), "white").save(scanned, "PDF")
    client = Client(razao_social="Cliente executável OCR sintético")
    user = User(nome="Usuário executável", email="ocr-exe@example.invalid", senha_hash="not-a-secret")
    db.add_all([client, user])
    db.flush()
    db.add(Proposal(
        numero=717,
        revisao="00",
        client_id=client.id,
        user_id=user.id,
        origem="upload_externo",
        pdf_path=scanned.name,
    ))
    db.commit()

    monkeypatch.setattr(index_service.platform, "system", lambda: "Windows")
    monkeypatch.setattr(
        index_service,
        "local_tesseract_ocr",
        lambda _path: (_ for _ in ()).throw(LocalOCRUnavailable("ocr_tesseract_missing")),
    )
    reply = AssistantService(db, None, output_dir=tmp_path).handle_message(
        message="O que consta no documento 717 sobre calibração?",
        request_id="missing-tesseract-assistant-query",
    )

    assert "executável tesseract ausente" in reply.message.lower()
    assert "não posso confirmar" in reply.message.lower()
    assert "não encontrei evidência suficiente" not in reply.message.lower()


def test_index_records_tesseract_output_with_page_citation(db, tmp_path, monkeypatch):
    from PIL import Image

    import app.services.document_index_service as index_service

    scanned = tmp_path / "scan.pdf"
    Image.new("RGB", (400, 250), "white").save(scanned, "PDF")
    client = Client(razao_social="Cliente OCR Windows sintético")
    user = User(nome="Usuário OCR", email="ocr-win@example.invalid", senha_hash="not-a-secret")
    db.add_all([client, user])
    db.flush()
    proposal = Proposal(
        numero=715,
        revisao="00",
        client_id=client.id,
        user_id=user.id,
        origem="upload_externo",
        pdf_path=scanned.name,
    )
    db.add(proposal)
    db.commit()
    monkeypatch.setattr(index_service.platform, "system", lambda: "Windows")
    monkeypatch.setattr(index_service, "local_tesseract_ocr", lambda _path: ["Inspeção local página sintética"])

    result = reindex_registered_documents(db, output_dir=tmp_path)
    db.commit()

    search = search_document_index(db, query="inspeção local sintética")
    assert result.ocr_pages == 1
    assert search.evidence[0].page == 1
    index = db.query(index_service.DocumentTextIndex).filter_by(source_id=proposal.id).first()
    assert index is not None
    assert index.language == "por"
